"""Corpus and extraction tool for the review (one CLI, several subcommands).

  python corpus.py file FOLDER       file the PDFs in FOLDER whose names start with a paper ID (P###) or a
                                     screening record ID (R####) into PDF_Library, as <ID>_<Title>.pdf
  python corpus.py sync              rename and move library PDFs to match papers.csv (ID, title, decision)
  python corpus.py text              cache the text of every included PDF in PDF_Text (<ID>.txt, <ID>.layout.txt)
  python corpus.py find ID REGEX     show matches in a paper's text with page numbers  [-c N context lines]
  python corpus.py page ID PAGE      print one page of a paper's layout text  [-n LINES] [-s START_REGEX]
  python corpus.py new BATCH         create Extraction/batches/BATCH.xlsx for the papers assigned to BATCH
  python corpus.py check BATCH       validate Extraction/batches/BATCH.xlsx against fields.csv and papers.csv
  python corpus.py tracker           regenerate Paper_Tracker.md from papers.csv, the library and the batches

Corpus/papers.csv lists every report assessed on the full text. Included papers carry a paper ID (P###), used
everywhere in extraction; record_id (R####) links a paper to the screening sheet. Extraction/batches/*.xlsx hold
the extracted values.
"""
import argparse
import csv
import re
import shutil
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import paths  # noqa: E402

COLLECTIONS = ["01_Detection_Counting", "02_Crown_Segmentation", "03_3D_Structure_Tree_Attributes",
               "04_Species_Classification", "05_Health_Disease_Mortality", "06_General_Purpose_Models"]
EXCLUDED = "Excluded_Full_Text"
ID_RX = re.compile(r"^(P\d{3}|R\d{4})(?=[-_ .])")
BATCH_COLS = ["paper_id", "item", "field", "value", "unit", "page", "location", "evidence", "status", "note", "checked"]
STATUSES = ["reported", "not_reported", "approximate", "derived"]


# ----------------------------------------------------------------------------- registry

def load_papers():
    with open(paths.PAPERS_CSV, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return rows, list(rows[0].keys())


def save_papers(rows, cols):
    with open(paths.PAPERS_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)


def resolve(rows, key):
    """Find a paper by paper ID (P###) or screening record ID (R####)."""
    hit = [r for r in rows if key in (r["paper_id"], r["record_id"])]
    if not hit:
        sys.exit(f"unknown ID {key}")
    return hit[0]


def canonical_name(row):
    t = re.sub(r"[^A-Za-z0-9]+", "_", row["title"]).strip("_")
    if len(t) > 120:
        t = t[:120].rsplit("_", 1)[0]
    pid = row["paper_id"] or row["record_id"]
    return f"{pid}_{t}.pdf" if pid else f"{t}.pdf"


def folder_for(row):
    return EXCLUDED if row["ft_decision"] == "Exclude" else row["collection"]


def in_review(row):
    return row["ft_decision"] in ("Include", "Pending")


# ----------------------------------------------------------------------------- PDFs

def cmd_file(a):
    rows, cols = load_papers()
    source = Path(a.folder).expanduser()
    if not source.is_dir():
        sys.exit(f"{source} is not a folder")
    filed, skipped = [], []
    for pdf in sorted(source.rglob("*.pdf")):
        m = ID_RX.match(pdf.name)
        hit = [r for r in rows if m and m.group(1) in (r["paper_id"], r["record_id"])]
        if not hit:
            skipped.append(f"{pdf.name}: no known ID at the start of the name")
            continue
        row = hit[0]
        dest = Path(paths.PDF_LIBRARY) / folder_for(row) / canonical_name(row)
        if dest.exists():
            skipped.append(f"{pdf.name}: {dest.name} already in the library")
            continue
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(pdf), dest)
        row["pdf_file"] = str(dest.relative_to(paths.PDF_LIBRARY))
        filed.append(f"{row['paper_id'] or row['record_id']} -> {row['pdf_file']}")
    save_papers(rows, cols)
    print(f"filed {len(filed)} PDF(s)")
    for line in filed + skipped:
        print("  ", line)


def cmd_sync(_):
    rows, cols = load_papers()
    moved, missing = [], []
    for r in rows:
        if not r["pdf_file"]:
            continue
        src = Path(paths.PDF_LIBRARY) / r["pdf_file"]
        if not src.exists():
            missing.append(r["paper_id"] or r["record_id"] or r["title"][:40])
            continue
        dest = Path(paths.PDF_LIBRARY) / folder_for(r) / canonical_name(r)
        if dest != src:
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(src), dest)
            r["pdf_file"] = str(dest.relative_to(paths.PDF_LIBRARY))
            moved.append(f"{src.name} -> {r['pdf_file']}")
    known = {r["pdf_file"] for r in rows if r["pdf_file"]}
    orphans = [str(p.relative_to(paths.PDF_LIBRARY)) for d in COLLECTIONS + [EXCLUDED]
               for p in (Path(paths.PDF_LIBRARY) / d).glob("*.pdf") if str(p.relative_to(paths.PDF_LIBRARY)) not in known]
    save_papers(rows, cols)
    print(f"moved {len(moved)}; registry entries with a missing file: {missing or 'none'}; unregistered PDFs: {orphans or 'none'}")


def cmd_text(_):
    rows, _ = load_papers()
    Path(paths.PDF_TEXT).mkdir(parents=True, exist_ok=True)
    keep = {r["paper_id"] for r in rows if in_review(r) and r["pdf_file"]}
    stale = [p for p in Path(paths.PDF_TEXT).glob("*.txt") if p.name.split(".")[0] not in keep]
    for p in stale:
        p.unlink()
    made = 0
    for r in rows:
        if r["paper_id"] not in keep:
            continue
        pdf = Path(paths.PDF_LIBRARY) / r["pdf_file"]
        for suffix, extra in ((".txt", []), (".layout.txt", ["-layout"])):
            out = Path(paths.PDF_TEXT) / f"{r['paper_id']}{suffix}"
            if not out.exists() or out.stat().st_mtime < pdf.stat().st_mtime:
                subprocess.run(["pdftotext", *extra, str(pdf), str(out)], check=True)
                made += 1
    print(f"wrote {made} text file(s) in {paths.PDF_TEXT}; removed {len(stale)} no longer needed")


def pages(key, raw=False):
    pid = resolve(load_papers()[0], key)["paper_id"]
    p = Path(paths.PDF_TEXT) / f"{pid}{'.txt' if raw else '.layout.txt'}"
    if not p.exists():
        sys.exit(f"no cached text for {pid}; run: corpus.py text")
    return p.read_text(encoding="utf-8", errors="ignore").split("\f")


def cmd_find(a):
    rx = re.compile(a.regex, re.I)
    for i, pg in enumerate(pages(a.id, a.raw), start=1):
        lines = pg.split("\n")
        for j, line in enumerate(lines):
            if rx.search(line):
                lo, hi = max(0, j - a.c), min(len(lines), j + a.c + 1)
                print(f"[p.{i}] " + " | ".join(x.strip() for x in lines[lo:hi] if x.strip())[:400])


def cmd_page(a):
    lines = pages(a.id)[a.page - 1].split("\n")
    start = next((j for j, x in enumerate(lines) if a.s and re.search(a.s, x, re.I)), 0)
    print("\n".join(lines[start:start + a.n]))


# ----------------------------------------------------------------------------- batches

def load_fields():
    with open(paths.FIELDS_CSV, encoding="utf-8") as f:
        return {r["field"]: r for r in csv.DictReader(f)}


def read_batch(path):
    ws = load_workbook(path, read_only=True, data_only=True)["values"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) if h is not None else "" for h in rows[0]]
    out = [{h: ("" if v is None else str(v).strip()) for h, v in zip(header, r)} for r in rows[1:] if any(v not in (None, "") for v in r)]
    return header, out


def load_batches():
    out = []
    for p in sorted(Path(paths.BATCHES).glob("B*.xlsx")):
        out += [dict(r, _batch=p.stem) for r in read_batch(p)[1]]
    return out


def cmd_new(a):
    rows, _ = load_papers()
    papers = [r for r in rows if r["batch"] == a.batch and in_review(r)]
    if not papers:
        sys.exit(f"no papers assigned to {a.batch}")
    path = Path(paths.BATCHES) / f"{a.batch}.xlsx"
    if path.exists():
        sys.exit(f"{path} exists; it holds extracted values and is never overwritten")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "values"
    ws.append(BATCH_COLS)
    lists = wb.create_sheet("lists")
    for i, v in enumerate(STATUSES, start=1):
        lists.cell(row=i, column=1, value=v)
    lists.sheet_state = "hidden"
    dv = DataValidation(type="list", formula1=f"lists!$A$1:$A${len(STATUSES)}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("I2:I5000")
    for col, w in zip("ABCDEFGHIJK", (8, 10, 22, 30, 8, 6, 14, 60, 13, 40, 18)):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:K5000"
    ps = wb.create_sheet("papers")
    ps.append(["paper_id", "record_id", "title", "year", "collection", "pdf_file"])
    for r in papers:
        ps.append([r["paper_id"], r["record_id"], r["title"], r["year"], r["collection"], r["pdf_file"]])
    for col, w in zip("ABCDEF", (9, 9, 90, 6, 30, 90)):
        ps.column_dimensions[col].width = w
    for c in ps[1]:
        c.font = Font(bold=True)
    for row in ps.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    print(f"created {path} for {len(papers)} papers")


def load_dataset_keys():
    if not paths.DATASETS_CSV.exists():
        return None
    with open(paths.DATASETS_CSV, encoding="utf-8") as f:
        return {r["dataset_key"] for r in csv.DictReader(f)}


def check_items(rows, fields):
    """Checks that need whole items: required fields per item, cross-references and the protocol 6 rules."""
    issues = []
    items, seen = {}, Counter()
    for r in rows:
        seen[(r["paper_id"], r["item"], r["field"])] += 1
        if r["item"]:
            items.setdefault((r["paper_id"], r["item"]), {})[r["field"]] = r
    issues += [f"{p} {i or '-'} {f}: recorded {n} times" for (p, i, f), n in seen.items() if n > 1]
    need = {}
    for f, d in fields.items():
        if d["item_prefix"] and d["required"] == "yes":
            need.setdefault(d["item_prefix"], []).append(f)
    keys = load_dataset_keys()
    refs = {"result_model": "model", "result_dataset": "dataset", "transfer_model": "model", "label_model": "model"}
    for (pid, item), f in items.items():
        prefix = re.sub(r"\d+$", "", item)
        missing = [x for x in need.get(prefix, []) if x not in f]
        if missing:
            issues.append(f"{pid} {item}: missing {missing} (add a not_reported row if absent)")
        for fld, pre in refs.items():
            v = f.get(fld, {}).get("value", "")
            if fld in f and f[fld]["status"] != "not_reported" and not (re.fullmatch(pre + r"\d+", v) and (pid, v) in items):
                issues.append(f"{pid} {item}: {fld}={v!r} is not a {pre} item of this paper")
        k = f.get("dataset_key")
        if k and k["status"] != "not_reported":
            if not re.fullmatch(r"[a-z0-9][a-z0-9_]*", k["value"]):
                issues.append(f"{pid} {item}: dataset_key {k['value']!r} must be lower case letters, digits and _")
            elif keys is not None and k["value"] not in keys:
                issues.append(f"{pid} {item}: dataset_key {k['value']!r} is not in datasets.csv")
        if prefix == "cf" and "cf_rule" in f and not f["cf_rule"]["note"]:
            issues.append(f"{pid} {item}: cf_rule needs the reason in note")
        if prefix == "tr":
            a_, b_ = f.get("value_in_domain"), f.get("value_out_of_domain")
            if a_ and b_ and "not_reported" not in (a_["status"], b_["status"]) and a_["unit"] != b_["unit"]:
                issues.append(f"{pid} {item}: in-domain and out-of-domain units differ ({a_['unit']} vs {b_['unit']})")
    by_paper = {}
    for (pid, item), f in items.items():
        by_paper.setdefault(pid, []).append((item, f))
    for pid, its in sorted(by_paper.items()):
        res = [(i, f) for i, f in its if i.startswith("res")]
        heads = [i for i, f in res if f.get("headline", {}).get("value") == "yes"]
        if len(heads) > 1:
            issues.append(f"{pid}: {len(heads)} headline results (at most one: the authors' main result)")
        for i, f in res:
            if f.get("result_task", {}).get("value") == "attribute_estimation" and "attribute" not in f:
                issues.append(f"{pid} {i}: attribute is required when result_task is attribute_estimation")
        prim = [(i, f) for i, f in res if f.get("primary", {}).get("value") == "yes"]
        uniq, setting = Counter(), {}
        for i, f in prim:
            g = lambda x: f.get(x, {}).get("value", "")  # noqa: E731
            if f.get("value", {}).get("status") == "approximate":
                issues.append(f"{pid} {i}: approximate values are never primary (rule V6)")
            if g("metric") == "other":
                issues.append(f"{pid} {i}: metric 'other' is never primary (protocol 6.3)")
            uniq[(g("result_model"), g("result_task"), g("attribute"), g("level"), g("metric"))] += 1
            setting.setdefault((g("result_task"), g("attribute"), g("level")), set()).add((g("result_dataset"), g("split"), g("aggregation")))
        issues += [f"{pid}: {n} primary results for model, task, attribute, level and metric {k} (one expected)" for k, n in uniq.items() if n > 1]
        issues += [f"{pid}: primary results for {k} mix dataset, split or aggregation {sorted(v)} (one setting expected)"
                   for k, v in setting.items() if len(v) > 1]
    return issues


def cmd_check(a):
    path = Path(paths.BATCHES) / f"{a.batch}.xlsx"
    header, rows = read_batch(path)
    fields = load_fields()
    papers = {r["paper_id"]: r for r in load_papers()[0] if r["paper_id"]}
    issues = []
    if header[:len(BATCH_COLS)] != BATCH_COLS:
        issues.append(f"columns must be exactly: {','.join(BATCH_COLS)}")
    for n, r in enumerate(rows, start=2):
        where = f"row {n} ({r.get('paper_id')}, {r.get('item') or '-'}, {r.get('field')})"
        p, fd = papers.get(r["paper_id"]), fields.get(r["field"])
        if not p:
            issues.append(f"{where}: unknown paper_id")
        elif not in_review(p):
            issues.append(f"{where}: paper is not in the review (ft_decision={p['ft_decision']})")
        elif p["batch"] != a.batch:
            issues.append(f"{where}: paper belongs to batch {p['batch']}")
        if not fd:
            issues.append(f"{where}: unknown field")
            continue
        if fd["item_prefix"] and not re.fullmatch(fd["item_prefix"] + r"\d+", r["item"]):
            issues.append(f"{where}: item must look like {fd['item_prefix']}1")
        if not fd["item_prefix"] and r["item"]:
            issues.append(f"{where}: single-value field must have an empty item")
        if r["status"] not in STATUSES:
            issues.append(f"{where}: status must be one of {STATUSES}")
        if r["status"] == "not_reported":
            continue
        if fd["type"] == "category":
            bad = [v for v in r["value"].split(";") if v.strip() not in fd["allowed_values"].split("|")]
            if bad:
                issues.append(f"{where}: {bad} not in {fd['allowed_values']}")
        if fd["type"] == "number":
            try:
                float(r["value"].replace(",", ""))
            except ValueError:
                issues.append(f"{where}: value is not a number")
        if "% | fraction" in fd["unit"] and not r["unit"]:
            issues.append(f"{where}: unit must be %, fraction or the physical unit")
        if r["status"] in ("reported", "approximate") and not (r["page"] and (r["location"] or r["evidence"])):
            issues.append(f"{where}: page and location or evidence are required")
        if r["status"] == "derived" and not r["note"]:
            issues.append(f"{where}: derived values need the formula in note")
        if r["checked"] and not (r["checked"] == "ok" or r["checked"].startswith("fix:")):
            issues.append(f"{where}: checked must be empty, 'ok' or 'fix: <value>'")
    for pid in sorted({r["paper_id"] for r in rows}):
        have = {r["field"] for r in rows if r["paper_id"] == pid}
        need = [f for f, d in fields.items() if d["required"] == "yes" and not d["item_prefix"] and f not in have]
        if need:
            issues.append(f"{pid}: required fields missing (add a not_reported row if absent): {need}")
    issues += check_items(rows, fields)
    expected = {pid for pid, p in papers.items() if p["batch"] == a.batch and in_review(p)}
    absent = sorted(expected - {r["paper_id"] for r in rows})
    if absent:
        issues.append(f"papers of {a.batch} with no values yet: {absent}")
    print(f"{path.name}: {len(rows)} values, {len({r['paper_id'] for r in rows})} papers, {len(issues)} issue(s)")
    for i in issues:
        print("  ", i)


# ----------------------------------------------------------------------------- tracker

HEADER = """# Paper Tracker: Individual-Tree Extraction Review

Generated by `Scripts/extraction/corpus.py tracker` from `Literature_Extraction/Corpus/papers.csv`, the PDF library and
the extraction batches. Do not edit by hand; decisions change in `papers.csv`, and the PDF and extraction columns update
themselves when files appear.

## Scope

A paper is in scope when the unit of analysis is the individual tree: detecting, counting or locating trees; delineating
crowns or segmenting trees (2D or 3D); segmenting tree components; classifying individual trees (species, health,
disease, mortality); or estimating attributes of individual trees. It must train or test a transformer-based model, use
remote sensing or close-range tree imagery, report quantitative results, and be published 2017 to 2025. Area-level
mapping (forest, tree or mangrove cover; deforestation and disturbance maps; wall-to-wall canopy height or biomass maps;
pixel-wise species maps) is out of scope.
"""

WHERE = """## Where things are

| What | Location |
|---|---|
| Registry of papers and decisions | `Literature_Extraction/Corpus/papers.csv` |
| PDFs, arranged as the Zotero collections | `Literature_Extraction/PDF_Library/<collection>/<ID>_<Title>.pdf` |
| Adding PDFs: name each file starting with its ID (`P###` or `R####`), put them in any folder, run `corpus.py file <folder>` | `Scripts/extraction/corpus.py` |
| Extraction protocol and field list | `Literature_Extraction/Extraction/protocol.md`, `fields.csv` |
| Extracted values, one workbook per batch | `Literature_Extraction/Extraction/batches/B##.xlsx` |
| Search and screening records (PRISMA) | `Literature_Extraction/Search/` |

IDs: `P###` is the paper ID used in the library, the extraction and the analysis. `R####` is the record ID of the same
paper in `Literature_Extraction/Search/screening/screening_sheet.xlsx`; papers found by other methods have none.

## Zotero structure (mirrored by the PDF library)

```
IVC Review (individual-tree extraction)
  01 Detection & Counting              PDF_Library/01_Detection_Counting
  02 Crown Segmentation                PDF_Library/02_Crown_Segmentation
  03 3D Structure & Tree Attributes    PDF_Library/03_3D_Structure_Tree_Attributes
  04 Species Classification            PDF_Library/04_Species_Classification
  05 Health, Disease & Mortality       PDF_Library/05_Health_Disease_Mortality
  06 General-Purpose Models            PDF_Library/06_General_Purpose_Models
  Excluded at full text                PDF_Library/Excluded_Full_Text
Related work (cited, not reviewed)     PDF_Library/Related_Work/{Reviews,Datasets,Architectures,Policy}
```
"""

TITLES = {"01_Detection_Counting": "Detection & Counting", "02_Crown_Segmentation": "Crown Segmentation",
          "03_3D_Structure_Tree_Attributes": "3D Structure & Tree Attributes",
          "04_Species_Classification": "Species Classification",
          "05_Health_Disease_Mortality": "Health, Disease & Mortality", "06_General_Purpose_Models": "General-Purpose Models"}


def cell(x):
    return str(x or "").replace("|", "\\|").replace("\n", " ").strip()


def cmd_tracker(_):
    rows, _ = load_papers()
    values = load_batches()
    extracted = {r["paper_id"] for r in values}
    checked = {pid for pid in extracted if all(v["checked"] for v in values if v["paper_id"] == pid)}
    has_pdf = {r["paper_id"] for r in rows if r["paper_id"] and r["pdf_file"] and (Path(paths.PDF_LIBRARY) / r["pdf_file"]).exists()}
    inc = sorted([r for r in rows if in_review(r)], key=lambda r: r["paper_id"])
    exc = [r for r in rows if r["ft_decision"] == "Exclude"]
    groups = Counter(r["study_group"] for r in inc if r["study_group"])
    studies = len(inc) - sum(n - 1 for n in groups.values())
    need = [r for r in inc if r["paper_id"] not in has_pdf]
    batches = sorted({r["batch"] for r in inc})
    link = lambda r: f"[doi](https://doi.org/{r['doi']})" if r["doi"] else ""
    mark = lambda ok: "✓" if ok else "–"

    L = [HEADER, "## Summary", ""]
    L.append(f"- **Included: {len(inc)} papers ({studies} studies)**; {sum(1 for r in inc if r['origin'] == 'database')} "
             f"found by the database searches and {sum(1 for r in inc if r['origin'] == 'other methods')} by other methods. "
             "The final corpus is confirmed after the extraction is cross-validated.")
    L.append(f"- Excluded at full text: {len(exc)} (" + "; ".join(
        f"{k}: {v}" for k, v in Counter(r["ft_reason"].split(' (')[0] for r in exc).most_common()) + ").")
    L.append(f"- PDFs in the library: {len(inc) - len(need)} of {len(inc)}. Extracted: {len(extracted)}; "
             f"cross-validated by you: {len(checked)}.")
    L.append("- Per collection: " + "; ".join(f"{TITLES[c]}: {sum(1 for r in inc if r['collection'] == c)}" for c in COLLECTIONS) + ".")
    pend = [r for r in inc if r["ft_decision"] == "Pending"]
    if pend:
        L.append(f"- Awaiting your decision: {len(pend)} ({', '.join(r['paper_id'] for r in pend)}); still counted as included above.")
        L += ["", f"## Decisions awaiting you ({len(pend)})", "", "| ID | Title | Batch | Proposal |", "|---|---|---|---|"]
        L += [f"| {r['paper_id']} | {cell(r['title'])} | {r['batch']} | {cell(r['ft_reason'])} |" for r in pend]
    L += ["", WHERE, "## Batch plan", "",
          "| Batch | Papers | PDFs in hand | Status | Extracted | Cross-validated |", "|---|---|---|---|---|---|"]
    for b in batches:
        bp = [r for r in inc if r["batch"] == b]
        got = sum(1 for r in bp if r["paper_id"] in has_pdf)
        status = "ready" if got == len(bp) else f"needs {len(bp) - got} PDF(s)"
        L.append(f"| {b} | {len(bp)} | {got} | {status} | {sum(1 for r in bp if r['paper_id'] in extracted)} | "
                 f"{sum(1 for r in bp if r['paper_id'] in checked)} |")
    L += ["", f"## PDFs still needed ({len(need)})", "",
          "Name each file starting with its ID, for example `P012_...pdf` or `R1416_...pdf`.", "",
          "| ID | Record | Title | Year | Venue | Link | Open access | Batch |", "|---|---|---|---|---|---|---|---|"]
    for r in sorted(need, key=lambda r: (r["batch"], r["paper_id"])):
        L.append(f"| {r['paper_id']} | {r['record_id']} | {cell(r['title'])} | {r['year']} | {cell(r['venue'][:40])} | "
                 f"{link(r)} | {r['open_access']} | {r['batch']} |")
    L += ["", f"## Included papers ({len(inc)})", ""]
    for c in COLLECTIONS:
        sec = [r for r in inc if r["collection"] == c]
        L += [f"### {TITLES[c]} ({len(sec)})", "",
              "| ID | Record | Title | Year | Venue | Link | Full text checked | PDF | Extracted | Batch |",
              "|---|---|---|---|---|---|---|---|---|---|"]
        for r in sec:
            ext = "cross-validated" if r["paper_id"] in checked else mark(r["paper_id"] in extracted)
            L.append(f"| {r['paper_id']} | {r['record_id'] or 'other methods'} | {cell(r['title'])} | {r['year']} | "
                     f"{cell(r['venue'][:40])} | {link(r)} | {'yes' if r['ft_checked'] == 'yes' else 'at extraction'} | "
                     f"{mark(r['paper_id'] in has_pdf)} | {ext} | {r['batch']} |")
        L.append("")
    L += [f"## Excluded at full text ({len(exc)})", "", "| Record | Title | Year | Reason |", "|---|---|---|---|"]
    for r in sorted(exc, key=lambda r: (r["record_id"] or "~", r["title"])):
        L.append(f"| {r['record_id'] or 'other methods'} | {cell(r['title'])} | {r['year']} | {cell(r['ft_reason'])} |")
    L.append("")
    out = paths.SURVEY / "Paper_Tracker.md"
    out.write_text("\n".join(L), encoding="utf-8")
    print(f"wrote {out}: {len(inc)} included ({studies} studies), {len(need)} PDFs still needed, batches {batches}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("sync", "text", "tracker"):
        sub.add_parser(name)
    sub.add_parser("file").add_argument("folder")
    f = sub.add_parser("find")
    f.add_argument("id")
    f.add_argument("regex")
    f.add_argument("-c", type=int, default=0)
    f.add_argument("--raw", action="store_true")
    p = sub.add_parser("page")
    p.add_argument("id")
    p.add_argument("page", type=int)
    p.add_argument("-n", type=int, default=80)
    p.add_argument("-s", default=None)
    sub.add_parser("new").add_argument("batch")
    sub.add_parser("check").add_argument("batch")
    a = ap.parse_args()
    {"file": cmd_file, "sync": cmd_sync, "text": cmd_text, "find": cmd_find, "page": cmd_page,
     "new": cmd_new, "check": cmd_check, "tracker": cmd_tracker}[a.cmd](a)


if __name__ == "__main__":
    main()
