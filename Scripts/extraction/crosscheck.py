"""Author cross-check of the extraction: draw a reproducible random sample of recorded numbers, then summarise the checks.

  python crosscheck.py sample [--seed 20261014] [--force]   write Extraction/crosscheck/crosscheck_sample.xlsx
  python crosscheck.py summary                              count the checks, list every value not marked OK

The sample has one randomly drawn result from every included paper (a primary result where the paper has one) and
EXTRA further numbers drawn at random by kind (other results, transfer values, model sizes and costs, dataset sizes,
GSD and point density). Only values recorded as printed are drawn (status reported, including values read from images);
derived, approximate and not-reported values, and results a paper quotes from other work (copied_from), are left out.
"""
import argparse
import csv
import json
import random
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import paths  # noqa: E402

OUT_DIR = paths.EXTRACTION / "crosscheck"
SAMPLE = OUT_DIR / "crosscheck_sample.xlsx"
RESULT_JSON = OUT_DIR / "crosscheck_result.json"
EXTRA = {"result": 20, "transfer": 8, "model": 13, "dataset": 13, "study": 7}
CHECKS = ["OK", "Wrong", "Not found"]
METRIC = {"f1": "F1", "ap50": "AP50 (AP at IoU 0.5)", "ap50_95": "AP50:95 (COCO AP)", "map": "mAP (as printed)", "iou": "IoU",
          "miou": "mIoU", "oa": "overall accuracy", "kappa": "kappa", "precision": "precision", "recall": "recall", "r2": "R2",
          "rmse": "RMSE", "mae": "MAE", "mape": "MAPE", "other": "Score (as named in the paper)"}
COLS = [("No.", 6), ("Paper", 8), ("Title", 34), ("PDF", 30), ("Page", 6), ("Where", 18), ("What to check", 60),
        ("Recorded value", 12), ("Unit", 9), ("Quote to look for", 48), ("Correct?", 11), ("Correct value or comment", 30),
        ("Kind", 14), ("Batch", 7), ("Item", 8), ("Field", 16)]


def load():
    with open(paths.PAPERS_CSV, encoding="utf-8") as f:
        papers = {r["paper_id"]: r for r in csv.DictReader(f) if r["paper_id"] and r["ft_decision"] == "Include"}
    items = defaultdict(dict)  # (pid, item) -> field -> row
    batch_of = {}
    for book in sorted(paths.BATCHES.glob("B*.xlsx")):
        ws = load_workbook(book, read_only=True)["values"]
        it = ws.iter_rows(values_only=True)
        head = next(it)
        for r in it:
            if not r[0]:
                continue
            row = {k: ("" if v is None else str(v)) for k, v in zip(head, r)}
            items[(row["paper_id"], row["item"])][row["field"]] = row
            batch_of[row["paper_id"]] = book.stem
    return papers, items, batch_of


def name_of(items, pid, item, field):
    return items.get((pid, item), {}).get(field, {}).get("value", "") or item


def candidates(papers, items):
    """Every number that may be drawn, as (pid, kind, item, field, what, primary)."""
    out = []
    for (pid, item), f in items.items():
        if pid not in papers:
            continue

        def add(field, kind, what, primary=False):
            r = f.get(field)
            if r and r["status"] == "reported" and r["evidence"] and r["page"]:
                out.append((pid, kind, item, field, what, primary))
        if item.startswith("res") and not f.get("copied_from", {}).get("value"):
            g = lambda k: f.get(k, {}).get("value", "")  # noqa: E731
            split = g("split")
            what = (f"{METRIC.get(g('metric'), g('metric'))} of '{name_of(items, pid, g('result_model'), 'model_name')}' on "
                    f"'{name_of(items, pid, g('result_dataset'), 'dataset_name')}'. Setting: {g('result_setting') or 'not given'}"
                    + (f" (split: {split})" if split and split != "not_stated" else ""))
            add("value", "result", what, g("primary") == "yes")
        elif item.startswith("tr"):
            g = lambda k: f.get(k, {}).get("value", "")  # noqa: E731
            model = name_of(items, pid, g("transfer_model"), "model_name")
            add("value_in_domain", "transfer", f"In-domain {METRIC.get(g('transfer_metric'), g('transfer_metric'))} of '{model}' (source: {g('transfer_source')})")
            add("value_out_of_domain", "transfer", f"Out-of-domain {METRIC.get(g('transfer_metric'), g('transfer_metric'))} of '{model}' on {g('transfer_target')}")
        elif item.startswith("model"):
            model = f.get("model_name", {}).get("value", item)
            add("total_params", "model", f"Number of parameters of '{model}'")
            add("flops", "model", f"Computational cost (FLOPs) of '{model}'")
        elif item.startswith("dataset"):
            ds = f.get("dataset_name", {}).get("value", item)
            add("n_images", "dataset", f"Number of images (tiles, patches or plots) in '{ds}'")
            add("n_trees", "dataset", f"Number of trees (crowns or instances) in '{ds}'")
        elif item == "":
            add("gsd", "study", "Ground sampling distance (spatial resolution) of the imagery")
            add("point_density", "study", "Point density of the point clouds")
    return out


def draw(papers, items, seed):
    rng = random.Random(seed)
    cands = sorted(candidates(papers, items))
    by_paper = defaultdict(list)
    for c in cands:
        if c[1] == "result":
            by_paper[c[0]].append(c)
    chosen = []
    for pid in sorted(papers):
        res = by_paper.get(pid, [])
        prim = [c for c in res if c[5]]
        if prim or res:
            chosen.append(rng.choice(prim or res))
    taken = set(chosen)
    for kind, n in EXTRA.items():
        pool = [c for c in cands if c[1] == kind and c not in taken]
        pick = rng.sample(pool, min(n, len(pool)))
        if len(pick) < n:  # fill a short pool with further results
            rest = [c for c in cands if c[1] == "result" and c not in taken and c not in pick]
            pick += rng.sample(rest, n - len(pick))
        chosen += pick
        taken.update(pick)
    missing = sorted(set(papers) - {c[0] for c in chosen})
    return chosen, missing


def write(papers, items, batch_of, chosen, seed):
    wb = Workbook()
    ws = wb.active
    ws.title = "check"
    ws.append([c for c, _ in COLS])
    for i, (_, w) in enumerate(COLS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font = Font(bold=True)
        ws.cell(1, i).fill = PatternFill("solid", fgColor="DDEBF7")
    rows = []
    for pid, kind, item, field, what, primary in chosen:
        r = items[(pid, item)][field]
        rows.append((pid, int(float(r["page"])), item, field, kind, what, primary, r))
    rows.sort(key=lambda t: (t[0], t[1], t[2], t[3]))
    for n, (pid, page, item, field, kind, what, primary, r) in enumerate(rows, start=1):
        p = papers[pid]
        label = {"result": "Result (primary)" if primary else "Result", "transfer": "Transfer value", "model": "Model size or cost",
                 "dataset": "Dataset size", "study": "Data property"}[kind]
        ws.append([n, pid, p["title"][:90], p["pdf_file"], page, r["location"], what, r["value"], r["unit"], r["evidence"], None, None,
                   label, batch_of.get(pid, ""), item, field])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLS)).column_letter}{ws.max_row}"
    lists = wb.create_sheet("lists")
    for i, v in enumerate(CHECKS, start=1):
        lists.cell(i, 1, v)
    dv = DataValidation(type="list", formula1=f"=lists!$A$1:$A${len(CHECKS)}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{ws.max_row}")
    how = wb.create_sheet("how to check", 0)
    text = [
        "Cross-check of the full-text extraction",
        "",
        f"{len(rows)} numbers drawn at random (seed {seed}, drawn {date.today().isoformat()} by Scripts/extraction/crosscheck.py sample):",
        "one result from every included paper (a primary result where the paper has one) and further numbers drawn by kind.",
        "",
        "For each row of the sheet 'check':",
        "1. Open the PDF (column PDF, under Literature_Extraction/PDF_Library) at the page in column Page: the position of the page in the PDF file, the first page of the file being 1 (the number printed on the page, which some viewers show, can differ, for example in book chapters).",
        "2. Find the place in column Where; the quote in column 'Quote to look for' is copied from that place (for tables read from",
        "   the rendered page, the quote is the table row as printed).",
        "3. Mark 'Correct?':",
        "   OK        the recorded value and unit are printed there, for the model, metric, dataset and setting in 'What to check';",
        "   Wrong     anything differs (value, unit, model, row or column, setting); write what the paper prints in the last column;",
        "   Not found the value cannot be found at the place given.",
        "Rows are sorted by paper and page, so each PDF is opened once. You can stop at any time; the file keeps your marks.",
        "",
        "Rule: every row must be OK. A value marked Wrong or Not found is corrected in its script, the whole paper is checked",
        "again, the same kind of value is checked across its batch, and a new sample is drawn for that batch.",
        "When done: Scripts/.venv/bin/python Scripts/extraction/crosscheck.py summary",
    ]
    for i, t in enumerate(text, start=1):
        how.cell(i, 1, t)
    how["A1"].font = Font(bold=True, size=13)
    how.column_dimensions["A"].width = 130
    summ = wb.create_sheet("summary")
    last = ws.max_row
    for i, (label, formula) in enumerate([("Rows", f"=COUNTA(check!A2:A{last})"), ("OK", f'=COUNTIF(check!K2:K{last},"OK")'),
                                          ("Wrong", f'=COUNTIF(check!K2:K{last},"Wrong")'), ("Not found", f'=COUNTIF(check!K2:K{last},"Not found")'),
                                          ("Not yet checked", f'=COUNTBLANK(check!K2:K{last})')], start=1):
        summ.cell(i, 1, label).font = Font(bold=True)
        summ.cell(i, 2, formula)
    summ.column_dimensions["A"].width = 18
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(SAMPLE)
    return rows


def summary():
    ws = load_workbook(SAMPLE, read_only=True)["check"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    ix = {k: i for i, k in enumerate(head)}
    rows = [r for r in it if r[0] is not None]
    marks = Counter((r[ix["Correct?"]] or "").strip() or "not yet checked" for r in rows)
    bad = [r for r in rows if (r[ix["Correct?"]] or "").strip() in ("Wrong", "Not found")]
    by_kind = defaultdict(Counter)
    for r in rows:
        by_kind[r[ix["Kind"]]][(r[ix["Correct?"]] or "").strip() or "not yet checked"] += 1
    print(f"{len(rows)} rows: " + ", ".join(f"{k} {v}" for k, v in sorted(marks.items())))
    for k in sorted(by_kind):
        print(f"  {k}: " + ", ".join(f"{m} {n}" for m, n in sorted(by_kind[k].items())))
    for r in bad:
        print(f"  {r[ix['Correct?']].upper()}: No. {r[ix['No.']]} {r[ix['Paper']]} p.{r[ix['Page']]} {r[ix['Item']]} {r[ix['Field']]} = {r[ix['Recorded value']]}"
              f" | {r[ix['Correct value or comment']] or ''}")
    checked = len(rows) - marks.get("not yet checked", 0)
    ok = marks.get("OK", 0)
    res = {"date": date.today().isoformat(), "rows": len(rows), "checked": checked, "ok": ok, "wrong": marks.get("Wrong", 0),
           "not_found": marks.get("Not found", 0), "papers": len({r[ix['Paper']] for r in rows})}
    if checked == len(rows) and ok == len(rows):
        res["upper_95_error_rate"] = round(1 - 0.05 ** (1 / len(rows)), 4)  # one-sided bound when no error is found
        print(f"all {len(rows)} values correct; one-sided 95% upper bound on the error rate: {res['upper_95_error_rate']:.2%}")
    with open(paths.PAPERS_CSV, encoding="utf-8") as f:  # the part of the sample that belongs to the current corpus
        current = {r["paper_id"] for r in csv.DictReader(f) if r["paper_id"] and r["ft_decision"] == "Include"}
    cur = [r for r in rows if r[ix["Paper"]] in current]
    cur_ok = sum(1 for r in cur if (r[ix["Correct?"]] or "").strip() == "OK")
    res["current_corpus"] = {"papers_in_corpus": len(current), "rows": len(cur), "papers": len({r[ix['Paper']] for r in cur}), "ok": cur_ok}
    if cur and cur_ok == len(cur):
        res["current_corpus"]["upper_95_error_rate"] = round(1 - 0.05 ** (1 / len(cur)), 4)
    print(f"current corpus ({len(current)} papers): {len(cur)} sampled values from {res['current_corpus']['papers']} papers, {cur_ok} OK"
          + (f"; upper bound {res['current_corpus']['upper_95_error_rate']:.2%}" if "upper_95_error_rate" in res["current_corpus"] else ""))
    RESULT_JSON.write_text(json.dumps(res, indent=2))
    sys.exit(1 if bad else 0)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["sample", "summary"])
    ap.add_argument("--seed", type=int, default=20261014)
    ap.add_argument("--force", action="store_true", help="overwrite an existing sample")
    a = ap.parse_args()
    if a.command == "summary":
        return summary()
    if SAMPLE.exists() and not a.force:
        sys.exit(f"{SAMPLE} exists (it may hold your checks); use --force to draw a new sample")
    papers, items, batch_of = load()
    chosen, missing = draw(papers, items, a.seed)
    rows = write(papers, items, batch_of, chosen, a.seed)
    kinds = Counter(("Result (primary)" if t[6] else "Result") if t[4] == "result" else t[4] for t in rows)
    print(f"wrote {SAMPLE}: {len(rows)} values from {len({t[0] for t in rows})} papers; " + ", ".join(f"{k} {v}" for k, v in sorted(kinds.items())))
    if missing:
        print(f"papers without a drawable result: {missing}")


if __name__ == "__main__":
    main()
