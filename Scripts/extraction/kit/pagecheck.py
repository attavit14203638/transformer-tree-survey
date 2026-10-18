"""Page check of the evidence quotes: is each quote visible on its cited page?

  python pagecheck.py B## [B## ...]

The text layer of some PDFs repeats text on another page (a hidden copy of the next page, or overlaid text), so a quote
can be found in the text of a page that does not show it. For every quote of 25 characters or more that the text layer
holds on several pages, or on a page other than the cited one, the candidate pages are rendered at 200 dpi and read by
tesseract; the quote counts as visible on a page when at least 80 percent of its words and numbers are read there.
Rows whose cited page does not show the quote while exactly one other page does are listed as fixes (exit code 1);
rows with no clear page are listed for a look at the rendered pages. Values read from images are not checked here
(see ocrcheck.py), and quotes that are not verbatim text (table cell references of B01 to B03) are skipped.
"""
import glob
import os
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import paths  # noqa: E402

OCR = os.path.realpath("/tmp") + "/kit/ocr_pages"  # tesseract cannot open files through the /tmp symlink on macOS
IMAGE_LOC = ("(image)", "(bar labels)", "(rotated)")


def squash(s):
    s = s.replace("\u2212", "-").replace("\u2009", " ").replace("\u202f", " ").replace("\u00a0", " ")
    for a, b in (("\ufb01", "fi"), ("\ufb02", "fl"), ("\ufb00", "ff"), ("\ufb03", "ffi"), ("\u2019", "'"), ("\u2018", "'"),
                 ("\u201c", '"'), ("\u201d", '"')):
        s = s.replace(a, b)
    return re.sub(r"\s+", "", s).lower()


_text = {}


def text_pages(pid):
    if pid not in _text:
        raw = open(f"{paths.PDF_TEXT}/{pid}.txt", encoding="utf-8", errors="ignore").read().split("\f")
        lay = open(f"{paths.PDF_TEXT}/{pid}.layout.txt", encoding="utf-8", errors="ignore").read().split("\f")
        _text[pid] = ["|".join((squash(a), squash(re.sub(r"-\n", "", a)), squash(b), squash(re.sub(r"-\n", "", b))))
                      for a, b in zip(raw, lay)]
    return _text[pid]


def ocr(pid, page):
    out = f"{OCR}/{pid}_{page}.txt"
    if not os.path.exists(out):
        os.makedirs(OCR, exist_ok=True)
        pdf = glob.glob(f"{paths.PDF_LIBRARY}/*/{pid}_*.pdf")[0]
        base = f"{OCR}/{pid}_{page}"
        subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-r", "200", "-png", "-singlefile", pdf, base], check=True)
        txt = subprocess.run(["tesseract", base + ".png", "-", "--psm", "3"], capture_output=True).stdout.decode("utf-8", "ignore")
        Path(out).write_text(txt)
        os.remove(base + ".png")
    return Path(out).read_text()


def words(s):
    return re.findall(r"[a-z]{4,}|\d+(?:\.\d+)?", s.lower())


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    todo = []
    for batch in sys.argv[1:]:
        ws = load_workbook(paths.BATCHES / f"{batch}.xlsx", read_only=True)["values"]
        it = ws.iter_rows(values_only=True)
        ix = {k: i for i, k in enumerate(next(it))}
        for r in it:
            if not r[0] or (r[ix["status"]] or "") != "reported" or not r[ix["evidence"]] or not r[ix["page"]]:
                continue
            loc = str(r[ix["location"]] or "")
            if any(k in loc for k in IMAGE_LOC):
                continue
            pid, page, ev = r[0], int(float(r[ix["page"]])), str(r[ix["evidence"]])
            longest = max(re.split(r"\s*(?:\.\.\.|\u2026)\s*", ev), key=len)
            frag = squash(longest)
            if len(frag) < 25:
                continue
            hits = [i + 1 for i, t in enumerate(text_pages(pid)) if frag in t]
            if hits and (page not in hits or len(hits) > 1):
                todo.append((batch, pid, r[ix["item"]] or "-", r[ix["field"]], page, hits, loc, longest))
    print(f"{len(todo)} quotes found on several pages of the text layer or off their cited page; reading the candidate pages by OCR")
    verdict = Counter()
    report = defaultdict(list)
    for batch, pid, item, field, page, hits, loc, longest in todo:
        ws_ = words(longest)
        seen = [p for p in sorted(set(hits) | {page})
                if sum(w in set(words(ocr(pid, p))) for w in ws_) >= 0.8 * max(1, len(ws_))]
        v = "ok" if page in seen else ("fix" if len(seen) == 1 else "look")
        verdict[v] += 1
        if v != "ok":
            report[(batch, pid)].append(f"  {v.upper()}: {item} {field} p.{page} ({loc}) visible on {seen or 'no page'} | {longest[:80]}")
    print(", ".join(f"{k} {n}" for k, n in sorted(verdict.items())))
    for (batch, pid), lines in sorted(report.items()):
        print(f"{batch} {pid}:")
        print("\n".join(lines))
    sys.exit(1 if verdict.get("fix") or verdict.get("look") else 0)


if __name__ == "__main__":
    main()
