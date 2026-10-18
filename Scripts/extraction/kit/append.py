"""Write papers' rows into their batch workbook from their extraction scripts (idempotent).

  python append.py B04 P005 [P006 ...]    run Extraction/scripts/B04/P005.py ... and replace those papers' rows
  python append.py B04 --all              re-run every script of the batch
  python append.py B04 --remove P006      delete those papers' rows (only papers set to Exclude in papers.csv)

For B01 to B03 the protocol 6 tags come from Extraction/scripts/legacy_tags.py and are added after the script's rows.
The `checked` column is kept for rows whose value did not change; rows whose value changed lose their check mark and
are listed, so they can be checked again.
"""
import csv
import importlib.util
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import paths  # noqa: E402

sys.path.insert(0, str(paths.EXTRACTION_SCRIPTS))
import legacy_tags  # noqa: E402

COLS = ["paper_id", "item", "field", "value", "unit", "page", "location", "evidence", "status", "note", "checked"]


def load_rows(batch, pid):
    path = paths.EXTRACTION_SCRIPTS / batch / f"{pid}.py"
    spec = importlib.util.spec_from_file_location(pid, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    rows = [tuple("" if c is None else c for c in r) for r in mod.ROWS]
    bad = [r for r in rows if len(r) != 9]
    if bad:
        sys.exit(f"{pid}: rows must have 9 columns (item ... note); first bad row: {bad[0]}")
    return legacy_tags.apply(pid, rows)


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    batch = sys.argv[1]
    remove = sys.argv[2] == "--remove"
    if remove:
        pids = sys.argv[3:]
        with open(paths.PAPERS_CSV, encoding="utf-8") as f:
            decision = {r["paper_id"]: r["ft_decision"] for r in csv.DictReader(f) if r["paper_id"]}
        still_in = [pid for pid in pids if decision.get(pid) != "Exclude"]
        if not pids or still_in:
            sys.exit(f"--remove only deletes rows of papers set to Exclude in papers.csv; not excluded: {still_in or 'no ID given'}")
    else:
        pids = sorted(p.stem for p in (paths.EXTRACTION_SCRIPTS / batch).glob("P*.py")) if sys.argv[2] == "--all" else sys.argv[2:]
    book = paths.BATCHES / f"{batch}.xlsx"
    wb = load_workbook(book)
    ws = wb["values"]
    header = [c.value for c in ws[1]]
    if header[:len(COLS)] != COLS:
        sys.exit(f"{book.name}: unexpected columns {header}")
    old = [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    new = {pid: [] for pid in pids} if remove else {pid: load_rows(batch, pid) for pid in pids}
    checks = {(r[0], r[1] or "", r[2]): (str(r[3] if r[3] is not None else ""), r[10]) for r in old if r[0] in new and r[10]}
    kept = [r for r in old if r[0] not in new]
    lost = []
    for pid, rows in new.items():
        for r in rows:
            prev = checks.get((pid, r[0], r[1]))
            mark = None
            if prev:
                if prev[0] == str(r[2]):
                    mark = prev[1]
                else:
                    lost.append(f"{pid} {r[0] or '-'} {r[1]}: {prev[0]!r} -> {r[2]!r} (was checked {prev[1]!r})")
            kept.append([pid, *[None if c == "" else c for c in r], mark])
    kept.sort(key=lambda r: r[0])  # stable: papers in ID order, rows of a paper in script order
    ws.delete_rows(2, ws.max_row)
    for r in kept:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in (row[3], row[7], row[9]):
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(book)
    for pid, rows in new.items():
        print(f"{pid}: {sum(1 for r in old if r[0] == pid)} rows removed" if remove else f"{pid}: {len(rows)} rows written")
    print(f"{book.name} now has {ws.max_row - 1} rows")
    if lost:
        print("values that changed after being checked (check them again):")
        for x in lost:
            print("  ", x)


if __name__ == "__main__":
    main()
