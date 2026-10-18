"""Self-check of a batch workbook: cross-references, duplicates, page ranges, and every number taken from the text
present on its cited page.  usage: verify.py B##"""
import collections
import re
import sys

from openpyxl import load_workbook

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parents[2]))
import paths  # noqa: E402

BATCH = sys.argv[1] if len(sys.argv) > 1 else sys.exit("usage: verify.py B##")
ws = load_workbook(f"{paths.BATCHES}/{BATCH}.xlsx", read_only=True)["values"]
rows = list(ws.iter_rows(values_only=True))
H = rows[0]
V = [dict(zip(H, ["" if c is None else str(c) for c in r])) for r in rows[1:] if r[0]]
by_pid = collections.defaultdict(list)
for r in V:
    by_pid[r["paper_id"]].append(r)
issues = collections.defaultdict(list)
NUMERIC = {"value", "value_in_domain", "value_out_of_domain", "label_value", "label_amount", "n_images", "n_trees", "total_params",
           "trainable_params", "flops", "gsd", "point_density", "cf_alternative"}
IMAGE_LOC = ("image", "bar labels", "rotated")
for pid, rs in by_pid.items():
    pages = open(f"{paths.PDF_TEXT}/{pid}.txt", encoding="utf-8", errors="ignore").read().split("\f")
    lay = open(f"{paths.PDF_TEXT}/{pid}.layout.txt", encoding="utf-8", errors="ignore").read().split("\f")
    dup = [k for k, n in collections.Counter((r["item"], r["field"]) for r in rs).items() if n > 1]
    if dup:
        issues[pid].append(f"duplicate (item, field): {dup}")
    items = collections.defaultdict(dict)
    for r in rs:
        if r["item"]:
            items[r["item"]][r["field"]] = r["value"]
    models = {i for i in items if i.startswith("model")}
    dsets = {i for i in items if i.startswith("dataset")}
    for i, f in items.items():
        if i.startswith("res"):
            for need in ("metric", "result_model", "result_dataset", "value"):
                if need not in f:
                    issues[pid].append(f"{i}: missing {need}")
            if f.get("result_model") not in models:
                issues[pid].append(f"{i}: result_model {f.get('result_model')} undefined")
            if f.get("result_dataset") not in dsets:
                issues[pid].append(f"{i}: result_dataset {f.get('result_dataset')} undefined")
        if i.startswith("tr") and f.get("transfer_model") not in models:
            issues[pid].append(f"{i}: transfer_model {f.get('transfer_model')} undefined")
        if i.startswith("de") and f.get("label_model") not in models:
            issues[pid].append(f"{i}: label_model {f.get('label_model')} undefined")
        if i.startswith("model"):
            for need in ("model_name", "model_role", "architecture_category", "backbone", "pretraining", "training_strategy"):
                if need not in f and not any(r["item"] == i and r["field"] == need for r in rs):
                    issues[pid].append(f"{i}: missing {need}")
    for r in rs:
        if r["page"]:
            p = int(float(r["page"]))
            if not 1 <= p <= len(pages):
                issues[pid].append(f"{r['item']} {r['field']}: page {p} outside 1..{len(pages)}")
                continue
            if r["field"] in NUMERIC and r["status"] == "reported" and not any(k in r["location"] for k in IMAGE_LOC):
                v = r["value"].replace(",", "")
                raw = (pages[p - 1] + "\n" + lay[p - 1]).replace("−", "-")
                # thousands separators (7,824 or 110 937) are joined in one copy; lists such as {1,10,100} stay intact in the other
                joined = re.sub(r"(?<=\d)(?:,\s?|\s|\u2009|\u202f)(?=\d{3}\b)", "", raw)
                pat = r"(?<![\d.])" + re.escape(v) + r"(?![\d])"
                if not (re.search(pat, raw) or re.search(pat, joined)):
                    issues[pid].append(f"{r['item']} {r['field']}={r['value']} not found on p.{p} ({r['location']})")
total = sum(len(v) for v in issues.values())
print(f"{len(V)} rows, {len(by_pid)} papers, {total} issue(s)")
for pid in sorted(issues):
    for m in issues[pid]:
        print(f"  {pid}: {m}")
