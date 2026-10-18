"""Second-pass check of values read from images: every reported number whose location says 'image' (tables or chart
labels stored as pictures, scanned pages) must be found again by tesseract OCR on its page.

  python ocrcheck.py B##            check the whole batch
  python ocrcheck.py B## P###       check one paper

Pages are read at 300 and 500 dpi (and rotated by 90 and 270 degrees when the location says 'rotated'). A value found
only with its decimal point dropped (OCR often reads 96.43 as 9643) is listed as a weak match: look at the page to
confirm it. Values not found at all must be checked on the rendered page (kit/render.py) before the batch is reported.
"""
import glob
import os
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import paths  # noqa: E402

if len(sys.argv) < 2:
    sys.exit(__doc__)
BATCH = sys.argv[1]
ONLY = sys.argv[2] if len(sys.argv) > 2 else None
OCR = os.path.realpath("/tmp") + "/kit/ocr"  # tesseract cannot open files through the /tmp symlink on macOS
os.makedirs(OCR, exist_ok=True)
NUMF = {"value", "value_in_domain", "value_out_of_domain", "total_params", "trainable_params", "flops", "gsd", "point_density",
        "n_images", "n_trees", "label_value", "label_amount", "cf_alternative"}

ws = load_workbook(f"{paths.BATCHES}/{BATCH}.xlsx", read_only=True)["values"]
rows = list(ws.iter_rows(values_only=True))
I = {h: i for i, h in enumerate(rows[0])}
need = defaultdict(list)
for r in rows[1:]:
    loc = str(r[I["location"]] or "")
    if (r[I["field"]] in NUMF and ("image" in loc or "rotated" in loc) and r[I["status"]] == "reported"
            and re.fullmatch(r"-?[\d.,]+", str(r[I["value"]] or "")) and (ONLY is None or r[I["paper_id"]] == ONLY)):
        need[(r[I["paper_id"]], int(r[I["page"]]), "rotated" in loc)].append(r)


def ocr(pid, page, rotated):
    pdf = glob.glob(f"{paths.PDF_LIBRARY}/*/{pid}_*.pdf")[0]
    txt = ""
    for dpi in (300, 500):
        base = f"{OCR}/{pid}_{page}_{dpi}"
        subprocess.run(["pdftoppm", "-f", str(page), "-l", str(page), "-r", str(dpi), "-png", "-singlefile", pdf, base], check=True)
        srcs = [base + ".png"]
        if rotated:
            for rot in (90, 270):
                Image.open(base + ".png").rotate(rot, expand=True).save(f"{base}_r{rot}.png")
                srcs.append(f"{base}_r{rot}.png")
        for src in srcs:
            txt += " " + subprocess.run(["tesseract", src, "-", "--psm", "6"], capture_output=True).stdout.decode("utf-8", "ignore")
    return txt.replace(",", "")


miss, weak, n = [], [], 0
for (pid, page, rotated), rs in sorted(need.items()):
    txt = ocr(pid, page, rotated)
    toks = set(re.findall(r"\d+(?:\.\d+)?", txt))
    digits = {t.replace(".", "") for t in toks}
    for r in rs:
        n += 1
        v = str(r[I["value"]]).replace(",", "")
        alts = {v, v.rstrip("0").rstrip(".") if "." in v else v, v + "0"}
        if alts & toks:
            continue
        hit = (pid, page, r[I["item"]], r[I["field"]], v, str(r[I["evidence"]] or "")[:80])
        (weak if v.replace(".", "") in digits else miss).append(hit)
print(f"{BATCH}: {n} values read from images checked; {len(weak)} weak matches (decimal point lost); {len(miss)} not found")
for label, xs in (("weak", weak), ("NOT FOUND", miss)):
    for x in xs:
        print(f"  {label}: {x}")
