"""Screening pipeline for the systematic search (PRISMA 2020).

  python screening.py build    ingest the database exports (Scopus CSV, WoS tab-delimited or RIS) and
                               results/arxiv_records.csv, de-duplicate, flag the reports already found
                               by the preliminary scoping searches, and write screening/screening_sheet.xlsx
  python screening.py sample   write the author's blind 20% sample (no AI decisions shown)
  python screening.py apply    copy AI title/abstract decisions (ai_decisions.csv) into the sheet
  python screening.py kappa    agreement between AI decisions and the author's blind sample
  python screening.py prisma   PRISMA 2020 flow counts -> screening/prisma_counts.json
  python screening.py searchlog  when each search was run -> search_log.json (export files and arXiv run log)

Title/abstract screening is AI-assisted (TA_R1); the first author screens a blind random 20% sample
and checks every record the AI includes. Full-text decisions, the papers found by other
methods and study groups are recorded in Literature_Extraction/Corpus/papers.csv, which prisma reads.
"""
import csv
import difflib
import glob
import json
import os
import random
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import paths  # noqa: E402

HERE = str(paths.SEARCH)
EXPORTS = str(paths.SEARCH_EXPORTS)
OUTDIR = os.path.join(HERE, "screening")
SHEET = os.path.join(OUTDIR, "screening_sheet.xlsx")
AUTHOR_SAMPLE = os.path.join(OUTDIR, "author_sample_20pct.xlsx")
AI_DECISIONS = os.path.join(OUTDIR, "ai_decisions.csv")
SEED = 2026
SECOND_SCREEN_FRACTION = 0.20

TA_CHOICES = ["Include", "Exclude", "Unsure"]
EXCLUSION_REASONS = [
    "No transformer-based model trained or evaluated",
    "Not a tree task",
    "Area-level mapping, not individual trees",
    "Not remote sensing or close-range tree imagery",
    "No quantitative evaluation",
    "Review, editorial or abstract only",
    "Duplicate report of an included study",
    "Outside date range",
    "Not in English",
    "Full text not retrievable",
    "Retracted publication",
]
CRITERIA = [
    "ELIGIBILITY CRITERIA: INDIVIDUAL-TREE EXTRACTION. Include a record only if ALL five apply.",
    "",
    "1. Model: trains or tests at least one transformer-based model, i.e. an architecture with transformer (multi-head "
    "self-attention) blocks such as ViT, Swin, SegFormer, DETR or Mask2Former, or a transformer-based foundation or "
    "vision-language model (e.g. SAM, CLIP, Grounding DINO). Transformers used only as baselines count. Generic attention "
    "modules alone (SE, CBAM, coordinate attention) do not.",
    "2. Task: the unit of analysis is the individual tree. That is: detecting, counting or locating individual trees "
    "(including trunks and saplings); delineating or segmenting tree crowns or individual trees, in 2D or 3D; segmenting "
    "tree components (stem, branches, foliage, wood and leaf); classifying individual trees or crowns by species, health, "
    "disease or mortality; or estimating attributes of individual trees (height, DBH, crown size, volume, biomass). "
    "Orchard, plantation and urban trees count. General-purpose models qualify if evaluated on at least one such task.",
    "3. Data: remote sensing or close-range imagery of trees (UAV, aerial, satellite, LiDAR or point clouds, ground-level "
    "tree images).",
    "4. Results: reports a quantitative evaluation.",
    "5. Publication: English; published 2017-2025 (the online date counts); journal or conference paper, book chapter, or "
    "preprint with full methods.",
    "",
    "EXCLUDE (area-level mapping): forest, tree, mangrove, orchard or plantation cover or extent maps; forest type or stand "
    "maps; deforestation, forest change or disturbance maps; wall-to-wall canopy height, biomass or carbon maps; pixel-wise "
    "species maps produced like land-cover classification.",
    "EXCLUDE (not a tree task): crops; grassland, shrubs, wetlands and other non-tree vegetation; fruit-, flower- or "
    "leaf-level analysis; wildfire and burned areas; general land cover; reviews and surveys; abstracts or editorials; "
    "theory-only papers; retracted papers.",
    "",
    "Mark Unsure when the title and abstract do not settle a criterion; Unsure records get a full-text check.",
    "Fill 'decision' for every row; give an 'exclusion_reason' for every Exclude.",
]
COLUMNS = [
    "record_id", "sources", "in_scoping_searches", "title", "abstract", "year", "venue",
    "doc_type", "doi", "arxiv_id", "keywords",
    "R2_sample", "TA_R1", "TA_R2", "TA_final", "TA_notes", "notes",
]


def norm_title(t):
    return re.sub(r"[^a-z0-9]", "", (t or "").lower())[:150]


def norm_doi(d):
    d = (d or "").strip().lower()
    return re.sub(r"^(https?://(dx\.)?doi\.org/|doi:)", "", d)


# ----------------------------------------------------------------------------- ingestion

def read_scopus_csv(path):
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        yield {
            "source": "Scopus", "source_id": r.get("EID", ""), "title": r.get("Title", ""),
            "abstract": r.get("Abstract", ""), "year": r.get("Year", ""), "doi": r.get("DOI", ""),
            "venue": r.get("Source title", ""), "doc_type": r.get("Document Type", ""),
            "keywords": "; ".join(x for x in (r.get("Author Keywords", ""), r.get("Index Keywords", "")) if x),
            "arxiv_id": "",
        }


def read_wos_tab(path):
    rows = list(csv.reader(open(path, encoding="utf-8-sig"), delimiter="\t", quoting=csv.QUOTE_NONE))
    header = rows[0]
    for vals in rows[1:]:
        r = dict(zip(header, vals))
        if not r.get("TI"):
            continue
        yield {
            "source": "Web of Science", "source_id": r.get("UT", ""), "title": r.get("TI", ""),
            "abstract": r.get("AB", ""), "year": r.get("PY", ""), "doi": r.get("DI", ""),
            "venue": r.get("SO", ""), "doc_type": r.get("DT", ""),
            "keywords": "; ".join(x for x in (r.get("DE", ""), r.get("ID", "")) if x), "arxiv_id": "",
        }


def read_ris(path, source):
    rec, out = defaultdict(list), []
    for line in open(path, encoding="utf-8-sig"):
        m = re.match(r"^([A-Z][A-Z0-9])  - ?(.*)$", line.rstrip("\n"))
        if not m:
            continue
        tag, val = m.groups()
        if tag == "ER":
            out.append(rec)
            rec = defaultdict(list)
        else:
            rec[tag].append(val.strip())
    for r in out:
        first = lambda *tags: next((r[t][0] for t in tags if r.get(t)), "")
        yield {
            "source": source, "source_id": first("AN", "ID", "UR"), "title": first("TI", "T1"),
            "abstract": first("AB", "N2"), "year": first("PY", "Y1", "DA")[:4], "doi": first("DO"),
            "venue": first("T2", "JO", "JF", "SO"), "doc_type": first("TY"),
            "keywords": "; ".join(r.get("KW", [])), "arxiv_id": "",
        }


def read_arxiv(path):
    for r in csv.DictReader(open(path, encoding="utf-8")):
        yield {
            "source": "arXiv", "source_id": r["arxiv_id"], "title": r["title"], "abstract": r["abstract"],
            "year": r["published"][:4], "doi": r.get("doi", ""), "venue": r.get("journal_ref", "") or "arXiv",
            "doc_type": "Preprint", "keywords": "", "arxiv_id": r["arxiv_id"],
        }


def ingest():
    raw = []
    for path in sorted(glob.glob(os.path.join(EXPORTS, "*"))):
        name, ext = os.path.basename(path).lower(), os.path.splitext(path)[1].lower()
        if ext == ".csv" and "scopus" in name:
            raw += list(read_scopus_csv(path))
        elif ext == ".txt" and ("wos" in name or "savedrecs" in name):
            raw += list(read_wos_tab(path))
        elif ext == ".ris":
            raw += list(read_ris(path, "Scopus" if "scopus" in name else "Web of Science"))
    arxiv = os.path.join(HERE, "results", "arxiv_records.csv")
    if os.path.exists(arxiv):
        raw += list(read_arxiv(arxiv))
    return raw


def deduplicate(raw):
    groups, by_doi, by_title = [], {}, {}
    for r in raw:
        d, t = norm_doi(r["doi"]), norm_title(r["title"])
        if r["arxiv_id"] and not d:
            d = f"10.48550/arxiv.{r['arxiv_id']}"
        idx = by_doi.get(d) if d else None
        if idx is None:
            idx = by_title.get(t)
        if idx is None:
            idx = len(groups)
            groups.append([])
        groups[idx].append(r)
        if d:
            by_doi[d] = idx
        by_title[t] = idx
    # A preprint and its published version can differ slightly in title: merge an arXiv-only group
    # into a published group whose normalised title is near-identical.
    keys = [norm_title(g[0]["title"]) for g in groups]
    preprint_only = [all(r["source"] == "arXiv" for r in g) for g in groups]
    alive = [True] * len(groups)
    blocks = defaultdict(list)
    for i, k in enumerate(keys):
        blocks[k[:12]].append(i)
    for ids in blocks.values():
        for i in (i for i in ids if preprint_only[i]):
            j = next((j for j in ids if j != i and alive[j] and not preprint_only[j]
                      and difflib.SequenceMatcher(None, keys[i], keys[j]).ratio() >= 0.95), None)
            if j is not None:
                groups[j].extend(groups[i])
                alive[i] = False
    groups = [g for g, a in zip(groups, alive) if a]
    merged = []
    for g in groups:
        best = max(g, key=lambda r: len(r["abstract"] or ""))
        m = dict(best)
        m["sources"] = "; ".join(sorted({r["source"] for r in g}))
        m["doi"] = next((norm_doi(r["doi"]) for r in g if r["doi"]), "")
        m["arxiv_id"] = next((r["arxiv_id"] for r in g if r["arxiv_id"]), "")
        merged.append(m)
    return merged


def corpus_lookup():
    rows = list(csv.DictReader(open(os.path.join(HERE, "corpus_index.csv"), encoding="utf-8")))
    by = {}
    for r in rows:
        for t in (r["title"], r["alt_title"]):
            if t:
                by[("t", norm_title(t))] = r["citation_key"]
        if r["doi"]:
            by[("d", norm_doi(r["doi"]))] = r["citation_key"]
        if r["arxiv_id"]:
            by[("a", r["arxiv_id"])] = r["citation_key"]
    return by


def corpus_match(m, lookup):
    """Citation key of the scoping-search report a record belongs to: DOI, arXiv ID or exact title first,
    then a near-identical title (databases rewrite some titles, e.g. 'TransU-Net++' as 'plus plus')."""
    key = (lookup.get(("d", m["doi"])) or lookup.get(("a", m["arxiv_id"]))
           or lookup.get(("t", norm_title(m["title"]))))
    if key:
        return key
    t = norm_title(m["title"])
    for (kind, ct), ck in lookup.items():
        if kind == "t":
            sm = difflib.SequenceMatcher(None, t, ct)
            if sm.quick_ratio() >= 0.9 and sm.ratio() >= 0.9:
                return ck
    return ""


def write_lists(wb):
    """Hidden sheet holding the dropdown choices. Excel rejects literal list validations longer than 255
    characters and splits them on commas, so every dropdown refers to a range on this sheet instead."""
    ws = wb.create_sheet("lists")
    ranges = {}
    for j, (key, values) in enumerate((("ta", TA_CHOICES), ("reasons", EXCLUSION_REASONS)), start=1):
        for i, v in enumerate(values, start=1):
            ws.cell(row=i, column=j, value=v)
        letter = get_column_letter(j)
        ranges[key] = f"lists!${letter}$1:${letter}${len(values)}"
    ws.sheet_state = "hidden"
    return ranges


def dropdown(ws, source, col, n):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{n}")


def build():
    raw = ingest()
    if not raw:
        sys.exit("No records found in exports/ or results/arxiv_records.csv")
    merged = deduplicate(raw)
    lookup = corpus_lookup()
    for m in merged:
        m["in_scoping_searches"] = corpus_match(m, lookup)
    merged.sort(key=lambda m: norm_title(m["title"]))
    rng = random.Random(SEED)
    sample = set(rng.sample(range(len(merged)), max(1, round(SECOND_SCREEN_FRACTION * len(merged)))))
    os.makedirs(OUTDIR, exist_ok=True)
    if os.path.exists(SHEET):
        sys.exit(f"{SHEET} exists; move it away before rebuilding so no decisions are overwritten.")

    wb = Workbook()
    ws = wb.active
    ws.title = "screening"
    ws.append(COLUMNS)
    for i, m in enumerate(merged, start=1):
        ws.append([f"R{i:04d}", m["sources"], m["in_scoping_searches"], m["title"], m["abstract"], m["year"],
                   m["venue"], m["doc_type"], m["doi"], m["arxiv_id"], m["keywords"],
                   1 if (i - 1) in sample else 0, "", "", "", "", ""])
    n = len(merged) + 1
    lists = write_lists(wb)
    for col in ("M", "N", "O"):
        dropdown(ws, lists["ta"], col, n)
    widths = {"A": 8, "B": 14, "C": 18, "D": 60, "E": 90, "F": 6, "G": 25, "H": 12, "I": 22, "J": 12, "K": 30,
              "L": 6, "M": 10, "N": 10, "O": 10, "P": 25, "Q": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_col=5):
        for c in row[3:5]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    grey = PatternFill("solid", fgColor="EEEEEE")
    for row in ws.iter_rows(min_row=2):
        if row[11].value == 0:
            row[13].fill = grey
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:Q{n}"
    wb.save(SHEET)

    by_source = Counter(r["source"] for r in raw)
    print(f"raw records: {dict(by_source)} total={len(raw)}")
    print(f"unique after de-duplication: {len(merged)}; duplicates removed: {len(raw) - len(merged)}")
    print(f"reports of the scoping searches found: {sum(1 for m in merged if m['in_scoping_searches'])}")
    print(f"second-screen sample: {len(sample)} records (seed {SEED})")
    json.dump({"raw_by_source": by_source, "raw_total": len(raw), "unique": len(merged)},
              open(os.path.join(OUTDIR, "ingest_summary.json"), "w"), indent=2)
    found = {m["in_scoping_searches"] for m in merged if m["in_scoping_searches"]}
    missing = sorted({r["citation_key"] for r in csv.DictReader(open(os.path.join(HERE, "corpus_index.csv")))} - found)
    with open(os.path.join(OUTDIR, "corpus_not_retrieved.csv"), "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["citation_key"])
        w.writerows([[k] for k in missing])
    print(f"reports of the scoping searches not retrieved by the database searches: {len(missing)} "
          "(see screening/corpus_not_retrieved.csv; they are recorded in Corpus/papers.csv as 'other methods')")


# ----------------------------------------------------------------------------- analysis

def load_sheet():
    """The screening sheet; without it (the public repository holds no abstracts), the screening log written by prisma."""
    if not os.path.exists(SHEET):
        rows = list(csv.DictReader(open(os.path.join(OUTDIR, "screening_log.csv"), encoding="utf-8")))
        return [{**r, "R2_sample": int(r["R2_sample"] or 0)} for r in rows]
    ws = load_workbook(SHEET, data_only=True)["screening"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


def cohen_kappa(a, b):
    n = len(a)
    labels = sorted(set(a) | set(b))
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(l) / n) * (b.count(l) / n) for l in labels)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0, po


def sample():
    """Write the author's blind 20% sample: titles and abstracts only, no AI decisions."""
    if os.path.exists(AUTHOR_SAMPLE):
        sys.exit(f"{AUTHOR_SAMPLE} exists; move it away before rewriting so no decisions are overwritten.")
    rows = [r for r in load_sheet() if r["R2_sample"] == 1]
    wb = Workbook()
    ws = wb.active
    ws.title = "screening"
    cols = ["record_id", "title", "abstract", "year", "venue", "doc_type", "doi", "decision", "exclusion_reason", "note"]
    ws.append(cols)
    for r in rows:
        ws.append([r["record_id"], r["title"], r["abstract"], r["year"], r["venue"], r["doc_type"], r["doi"], "", "", ""])
    n = len(rows) + 1
    lists = write_lists(wb)
    dropdown(ws, lists["ta"], "H", n)
    dropdown(ws, lists["reasons"], "I", n)
    for col, w in {"A": 8, "B": 60, "C": 100, "D": 6, "E": 25, "F": 12, "G": 22, "H": 11, "I": 34, "J": 30}.items():
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_col=3):
        for c in row[1:3]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    crit = wb.create_sheet("criteria")
    for line in CRITERIA:
        crit.append([line])
    crit.column_dimensions["A"].width = 140
    wb.save(AUTHOR_SAMPLE)
    print(f"wrote {AUTHOR_SAMPLE} ({len(rows)} records)")


def kappa():
    rows = load_sheet()
    ai = {r["record_id"]: r["TA_R1"] for r in rows if r["TA_R1"]}
    if os.path.exists(AUTHOR_SAMPLE):
        ws = load_workbook(AUTHOR_SAMPLE, data_only=True)["screening"]
        header = [c.value for c in ws[1]]
        author = {d["record_id"]: d["decision"] for d in (dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2))
                  if d["decision"]}
    else:  # the public repository: the author's blind decisions are column TA_R2 of the screening log
        author = {r["record_id"]: r["TA_R2"] for r in rows if r["R2_sample"] == 1 and r["TA_R2"]}
    ids = sorted(set(ai) & set(author))
    binar = lambda v: "exclude" if v == "Exclude" else "include"
    k, po = cohen_kappa([binar(ai[i]) for i in ids], [binar(author[i]) for i in ids])
    disagree = [i for i in ids if binar(ai[i]) != binar(author[i])]
    out = {"n_double_screened": len(ids), "percent_agreement": round(100 * po, 1), "cohen_kappa": round(k, 3),
           "disagreements": disagree}
    json.dump(out, open(os.path.join(OUTDIR, "screening_agreement.json"), "w"), indent=2)
    print(out)


def apply_decisions():
    """Copy AI title/abstract decisions from ai_decisions.csv into TA_R1 / TA_notes of the sheet."""
    dec = {r["record_id"]: r for r in csv.DictReader(open(AI_DECISIONS, encoding="utf-8"))}
    wb = load_workbook(SHEET)
    ws = wb["screening"]
    header = [c.value for c in ws[1]]
    col_id, col_r1, col_notes = header.index("record_id"), header.index("TA_R1"), header.index("TA_notes")
    n = 0
    for row in ws.iter_rows(min_row=2):
        d = dec.get(row[col_id].value)
        if d:
            row[col_r1].value = d["decision"]
            row[col_notes].value = "; ".join(x for x in (d["reason"], d["note"]) if x)
            n += 1
    wb.save(SHEET)
    print(f"applied {n} AI decisions; {Counter(d['decision'] for d in dec.values())}")


def prisma():
    rows = load_sheet()
    ingest_summary = json.load(open(os.path.join(OUTDIR, "ingest_summary.json")))
    ta_final = lambda r: r["TA_final"] or r["TA_R1"]
    undecided = [r["record_id"] for r in rows if not ta_final(r)]
    if undecided:
        print(f"WARNING: {len(undecided)} records without a title/abstract decision")
    registry = list(csv.DictReader(open(paths.PAPERS_CSV, encoding="utf-8")))
    papers = {p["record_id"]: p for p in registry if p["record_id"]}
    reason = lambda p: p["ft_reason"].split(" (")[0]
    not_retrievable = lambda p: p["ft_decision"] == "Exclude" and reason(p) == "Full text not retrievable"
    sought = [papers[r["record_id"]] for r in rows if ta_final(r) in ("Include", "Unsure")]
    pending = [p["record_id"] for p in sought if p["ft_decision"] == "Pending"]
    if pending:
        print(f"WARNING: {len(pending)} reports still without a full-text decision: {pending}")
    not_retrieved = [p for p in sought if not_retrievable(p)]
    assessed = [p for p in sought if not not_retrievable(p) and p["ft_decision"] != "Pending"]
    excluded_ft = Counter(reason(p) for p in assessed if p["ft_decision"] == "Exclude")
    included_db = [p for p in assessed if p["ft_decision"] == "Include"]

    other = [p for p in registry if p["origin"] == "other methods"]
    other_by_method = Counter("preliminary scoping searches" for _ in other)
    other_assessed = [p for p in other if not not_retrievable(p)]
    other_included = [p for p in other_assessed if p["ft_decision"] == "Include"]
    other_excluded = Counter(reason(p) for p in other_assessed if p["ft_decision"] == "Exclude")

    reports = len(included_db) + len(other_included)
    groups = Counter(p["study_group"] for p in included_db + other_included if p["study_group"])
    merged_reports = sum(n - 1 for n in groups.values())
    counts = {
        "identification": {"databases": ingest_summary["raw_by_source"],
                           "records_total": ingest_summary["raw_total"],
                           "duplicates_removed": ingest_summary["raw_total"] - ingest_summary["unique"]},
        "screening": {"records_screened": len(rows),
                      "records_excluded": sum(1 for r in rows if ta_final(r) == "Exclude"),
                      "reports_sought": len(sought), "reports_not_retrieved": len(not_retrieved),
                      "reports_assessed": len(assessed), "reports_excluded": dict(excluded_ft)},
        "other_methods": {"records_by_method": dict(other_by_method), "reports_sought": len(other),
                          "reports_not_retrieved": len(other) - len(other_assessed),
                          "reports_assessed": len(other_assessed), "reports_excluded": dict(other_excluded)},
        "included": {"reports": reports, "studies": reports - merged_reports,
                     "via_databases": len(included_db), "via_other_methods": len(other_included)},
    }
    json.dump(counts, open(os.path.join(OUTDIR, "prisma_counts.json"), "w"), indent=2)
    public = ["record_id", "sources", "title", "year", "doi", "arxiv_id",
              "R2_sample", "TA_R1", "TA_R2", "TA_final", "ft_decision", "ft_reason"]
    with open(os.path.join(OUTDIR, "screening_log.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=public, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            p = papers.get(r["record_id"], {})
            w.writerow({**r, "ft_decision": p.get("ft_decision", ""), "ft_reason": p.get("ft_reason", "")})
    print(json.dumps(counts, indent=2))


def searchlog():
    """When each search was run: the arXiv run log, and the time each database export was saved. Times and the reported
    dates are in UTC."""
    def saved(pattern):
        t = min(os.path.getmtime(f) for f in glob.glob(os.path.join(EXPORTS, pattern)))
        return datetime.fromtimestamp(t, timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    arxiv = json.load(open(os.path.join(HERE, "results", "arxiv_run.json")))
    log = {"timezone": "UTC",
           "sources": {"Scopus": {"run_at_utc": saved("Scopus*.csv"), "record": "time the export file was saved"},
                       "Web of Science": {"run_at_utc": saved("wos_*.txt"), "record": "time the first export file was saved"},
                       "arXiv": {"run_at_utc": arxiv["run_at_utc"], "record": "results/arxiv_run.json"}}}
    json.dump(log, open(os.path.join(HERE, "search_log.json"), "w"), indent=2)
    print(json.dumps(log, indent=2))


if __name__ == "__main__":
    {"build": build, "sample": sample, "apply": apply_decisions, "kappa": kappa,
     "prisma": prisma, "searchlog": searchlog}[sys.argv[1] if len(sys.argv) > 1 else "build"]()
